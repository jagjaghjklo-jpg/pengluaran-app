import pandas as pd
import re

from openpyxl import load_workbook

FILE = "data/pengeluaran_2026.xlsx"

BULAN = [
    "Januari", "Februari", "Maret", "April", "Mei", "Juni",
    "Juli", "Agustus", "September", "Oktober", "November", "Desember"
]

KOLOM_LENGKAP = [
    "Tanggal", "Kategori", "Barang", "Harga", "Jumlah", "Subtotal", "Diskon"
]


def load_data():

    semua = []
    gagal = []

    match = re.search(r"(20\d{2})", FILE)
    tahun = int(match.group(1)) if match else None

    for bulan in BULAN:

        try:
            raw = pd.read_excel(FILE, sheet_name=bulan, header=None)

            if raw.dropna(how="all").empty:
                continue

            baris_header = None

            for i in range(len(raw)):
                nilai_baris = raw.iloc[i].astype(str).str.strip().tolist()
                if "Tanggal" in nilai_baris:
                    baris_header = i
                    break

            if baris_header is None:
                gagal.append(f"{bulan} gagal dibaca: header 'Tanggal' tidak ditemukan")
                continue

            # ambil nama-nama kolom asli persis dari baris header di Excel
            header_asli = (
                raw.iloc[baris_header]
                .astype(str)
                .str.strip()
                .tolist()
            )

            df = raw.iloc[baris_header + 1:].copy()
            df.columns = header_asli[:df.shape[1]]

            # cocokkan nama kolom asli ke nama standar (toleran variasi penulisan)
            def normalisasi(teks):
                return re.sub(r"[^a-z0-9]", "", str(teks).lower())

            mapping = {}
            for kolom_asli in df.columns:
                asli_bersih = normalisasi(kolom_asli)
                for kolom_standar in KOLOM_LENGKAP:
                    standar_bersih = normalisasi(kolom_standar)
                    if standar_bersih in asli_bersih:
                        mapping[kolom_asli] = kolom_standar
                        break

            df = df.rename(columns=mapping)

            # kolom standar yang gak ketemu di sheet ini diisi kosong
            for kolom in KOLOM_LENGKAP:
                if kolom not in df.columns:
                    df[kolom] = None

            df = df[KOLOM_LENGKAP]

            df["_excel_row"] = df.index + 1
            df["_sheet"] = bulan

            kolom_cek = ["Kategori", "Barang", "Harga", "Jumlah", "Subtotal"]
            df = df[df[kolom_cek].notna().any(axis=1)].copy()

            df["Tanggal"] = df["Tanggal"].ffill()
            df = df[df[kolom_cek].notna().any(axis=1)].copy()

            df["Bulan"] = bulan
            df["Tahun"] = tahun

            semua.append(df)

        except Exception as e:
            gagal.append(f"{bulan} gagal dibaca: {e}")

    if len(semua) == 0:
        return pd.DataFrame(), gagal

    hasil = pd.concat(semua, ignore_index=True)

    hasil["Harga"] = pd.to_numeric(hasil["Harga"], errors="coerce")
    hasil["Jumlah"] = pd.to_numeric(hasil["Jumlah"], errors="coerce")
    hasil["Subtotal"] = pd.to_numeric(hasil["Subtotal"], errors="coerce")
    hasil["Diskon"] = pd.to_numeric(hasil["Diskon"], errors="coerce")

    hasil["Kategori"] = hasil["Kategori"].astype(str).str.strip()
    hasil["Barang"] = hasil["Barang"].astype(str).str.strip()
    hasil["Tanggal"] = hasil["Tanggal"].astype(str).str.strip()

    hasil["Kategori"] = hasil["Kategori"].replace("nan", "")
    hasil["Barang"] = hasil["Barang"].replace("nan", "")

    # buang baris ringkasan/total yang ikut kebaca sebagai data
    hasil = hasil[
        ~hasil.apply(
            lambda row: row.astype(str).str.upper().str.contains("TOTAL").any(),
            axis=1
        )
    ].reset_index(drop=True)

    hasil = hasil[[
        "Tahun", "Bulan", "Tanggal", "Kategori", "Barang",
        "Harga", "Jumlah", "Subtotal", "Diskon", "_sheet", "_excel_row"
    ]]

    return hasil, gagal

def tambah_baris(bulan, tanggal, kategori, barang, harga, jumlah, subtotal):
    wb = load_workbook(FILE)
    ws = wb[bulan]

    baris_baru = ws.max_row + 1

    ws.cell(baris_baru, 1).value = tanggal
    ws.cell(baris_baru, 2).value = kategori
    ws.cell(baris_baru, 3).value = barang
    ws.cell(baris_baru, 4).value = harga
    ws.cell(baris_baru, 5).value = jumlah
    ws.cell(baris_baru, 6).value = subtotal
    ws.cell(baris_baru, 7).value = max((harga * jumlah) - subtotal, 0)
   
    wb.save(FILE)
    
def hapus_baris(daftar_hapus):
    wb = load_workbook(FILE)

    per_sheet = {}
    for item in daftar_hapus:
        sheet = item["_sheet"]
        baris = int(item["_excel_row"])
        per_sheet.setdefault(sheet, []).append(baris)

    for sheet_name, daftar_baris in per_sheet.items():
        ws = wb[sheet_name]
        for baris in sorted(daftar_baris, reverse=True):
            ws.delete_rows(baris, 1)

    wb.save(FILE)

def hapus_baris(daftar_hapus):
    wb = load_workbook(FILE)

    per_sheet = {}
    for item in daftar_hapus:
        sheet = item["_sheet"]
        baris = int(item["_excel_row"])
        per_sheet.setdefault(sheet, []).append(baris)

    for sheet_name, daftar_baris in per_sheet.items():
        ws = wb[sheet_name]
        for baris in sorted(daftar_baris, reverse=True):
            ws.delete_rows(baris, 1)

    wb.save(FILE)